# Kirill sinonimlar va tutuq/bo'shliq belgilari ataylab yozilgan (lotinga o'xshash harflar).
# ruff: noqa: RUF001
"""Tovarlarni Excel/CSV dan import qilish (T-051, `docs/specs/product-import.md`).

Uch bosqich: o'qish (`read_table` + `suggest_mapping`), tekshirish (`analyze`)
va qo'llash (`apply_import`, bitta tranzaksiya, bulk). Bulk amallar signal
bermaydi — pull uchun `ChangeLog` yozuvlari shu yerda qo'shiladi.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from django.db import transaction
from django.utils import timezone

from apps.core.exceptions import DomainError
from apps.core.ids import new_id
from apps.core.money import line_total
from apps.sync.models import ChangeLog, ChangeOp

from .barcodes import BarcodeError, normalize_barcode
from .models import Category, PriceHistory, Product, ProductBarcode, ProductImport, Unit

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20_000
SAMPLE_ROWS = 20
MAX_ERRORS = 500

FIELDS = ("name", "barcode", "category", "unit", "sell_price", "cost_price", "stock")
REQUIRED = ("name", "sell_price")

# Sarlavha → maydon (kichik harf, tutuq belgisi `'`, ortiqcha bo'shliqsiz).
HEADER_ALIASES = {
    "name": {"nomi", "nom", "name", "tovar", "tovar nomi", "название", "наименование", "товар"},
    "barcode": {
        "shtrix-kod",
        "shtrix kod",
        "shtrixkod",
        "barcode",
        "ean",
        "штрих-код",
        "штрихкод",
        "штрих код",
    },
    "category": {"kategoriya", "guruh", "category", "категория", "группа"},
    "unit": {"birlik", "o'lchov", "unit", "единица", "ед", "ед.", "ед. изм.", "бирлик"},
    "sell_price": {"narx", "sotish narxi", "price", "цена", "цена продажи", "нарх"},
    "cost_price": {
        "tannarx",
        "kirim narxi",
        "cost",
        "себестоимость",
        "закупочная цена",
        "таннарх",
    },
    "stock": {"qoldiq", "miqdor", "stock", "остаток", "количество", "қолдиқ"},
}

# Birlik matni → standart kod.
UNIT_ALIASES = {
    "pcs": {"pcs", "dona", "ta", "шт", "шт.", "штук", "дона"},
    "kg": {"kg", "кг", "kilogramm", "килограмм"},
    "g": {"g", "gr", "г", "гр", "gramm", "грамм"},
    "l": {"l", "litr", "л", "литр"},
    "m": {"m", "metr", "м", "метр"},
    "pack": {"pack", "quti", "pachka", "упак", "уп", "уп.", "коробка", "қути"},
}

_APOSTROPHES = re.compile("[ʻʼ‘’`]")
_SPACES = re.compile(r"[\s  ]+")


def _key(text: str) -> str:
    return _SPACES.sub(" ", _APOSTROPHES.sub("'", str(text))).strip().lower()


def _cell(value) -> str:
    """Katak → matn. Excel soni: butun bo'lsa — kasrsiz (shtrix-kod `4.78e12` emas)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(Decimal(repr(value)).normalize(), "f")
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value).strip()


def _read_xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise DomainError("invalid_file", "File is not a valid .xlsx workbook.") from exc
    try:
        sheet = workbook.worksheets[0]
        return [[_cell(v) for v in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_csv(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise DomainError("invalid_file", "CSV encoding is not supported.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


def read_table(data: bytes, filename: str) -> list[list[str]]:
    """Fayl → qatorlar (birinchisi — sarlavha); bo'sh qatorlar tashlanadi."""
    if len(data) > MAX_BYTES:
        raise DomainError("file_too_large", "File is larger than 10 MB.")
    name = filename.lower()
    if name.endswith(".csv"):
        rows = _read_csv(data)
    elif name.endswith(".xlsx") or data[:2] == b"PK":
        rows = _read_xlsx(data)
    else:
        raise DomainError("invalid_file", "Only .xlsx and .csv files are supported.")
    rows = [row for row in rows if any(cell for cell in row)]
    if len(rows) < 2:
        raise DomainError("invalid_file", "File has no data rows.")
    if len(rows) - 1 > MAX_ROWS:
        raise DomainError("too_many_rows", "File has more than 20 000 rows.")
    width = max(len(row) for row in rows)
    return [row + [""] * (width - len(row)) for row in rows]


def suggest_mapping(header: list[str]) -> dict[str, int | None]:
    keys = [_key(h) for h in header]
    mapping: dict[str, int | None] = {}
    for name in FIELDS:
        aliases = HEADER_ALIASES[name]
        mapping[name] = next((i for i, k in enumerate(keys) if k in aliases), None)
    return mapping


def clean_mapping(raw, width: int) -> dict[str, int | None]:
    if not isinstance(raw, dict):
        raise DomainError("invalid_mapping", "mapping must be an object.")
    mapping: dict[str, int | None] = {}
    for name in FIELDS:
        value = raw.get(name)
        if value is not None and (
            not isinstance(value, int) or isinstance(value, bool) or not 0 <= value < width
        ):
            raise DomainError("invalid_mapping", f"Column for {name} is out of range.")
        mapping[name] = value
    for name in REQUIRED:
        if mapping[name] is None:
            raise DomainError("invalid_mapping", f"Column for {name} is required.")
    return mapping


class BadNumber(ValueError):
    """Katakdagi matn son emas (yoki butun bo'lishi kerak edi)."""


def _decimal(text: str) -> Decimal:
    clean = _SPACES.sub("", text).replace(",", ".")
    if not re.fullmatch(r"-?\d+(\.\d+)?", clean):
        raise BadNumber(text)
    try:
        return Decimal(clean)
    except InvalidOperation as exc:
        raise BadNumber(text) from exc


def parse_money(text: str) -> int:
    """So'm → tiyin (half-up, 2 xona)."""
    return int((_decimal(text) * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))


def parse_qty(text: str, *, fractional: bool) -> int:
    value = _decimal(text) * 1000
    milli = value.quantize(Decimal(1), rounding=ROUND_HALF_UP)
    if not fractional and value % 1000:
        raise BadNumber(text)
    return int(milli)


@dataclass
class ParsedRow:
    row: int  # Excel qator raqami (sarlavha — 1)
    name: str
    barcode: str
    category: str
    unit: Unit
    sell: int
    cost: int
    stock: int
    product: Product | None = None


@dataclass
class Analysis:
    rows: list[ParsedRow] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)

    @property
    def creates(self) -> int:
        return sum(1 for r in self.rows if r.product is None)

    @property
    def updates(self) -> int:
        return sum(1 for r in self.rows if r.product is not None)

    def summary(self) -> dict:
        return {
            "creates": self.creates,
            "updates": self.updates,
            "error_count": len(self.errors),
            "errors": self.errors[:MAX_ERRORS],
        }


def _units(tenant_id) -> dict[str, Unit]:
    """Matn → birlik: kod, nom va standart sinonimlar."""
    by_text: dict[str, Unit] = {}
    for unit in Unit.objects.filter(tenant_id=tenant_id, deleted_at__isnull=True):
        by_text[_key(unit.code)] = unit
        if unit.name:
            by_text[_key(unit.name)] = unit
        for alias in UNIT_ALIASES.get(unit.code, ()):
            by_text.setdefault(alias, unit)
    return by_text


def analyze(product_import: ProductImport, mapping: dict[str, int | None]) -> Analysis:
    tenant_id = product_import.tenant_id
    units = _units(tenant_id)
    default_unit = units.get("pcs")
    existing = {
        b.barcode: b.product
        for b in ProductBarcode.objects.filter(
            tenant_id=tenant_id, deleted_at__isnull=True, product__deleted_at__isnull=True
        ).select_related("product", "product__unit")
    }
    result = Analysis()
    seen: set[str] = set()

    def value(row: list[str], name: str) -> str:
        index = mapping.get(name)
        return row[index].strip() if index is not None else ""

    for offset, row in enumerate(product_import.rows[1:]):
        number = offset + 2
        problems: list[tuple[str, str]] = []
        name = value(row, "name")
        if not name:
            problems.append(("name", "required"))
        elif len(name) > 200:
            problems.append(("name", "too_long"))

        barcode = value(row, "barcode")
        if barcode:
            try:
                barcode = normalize_barcode(barcode)
            except BarcodeError as exc:
                problems.append(("barcode", exc.code))
            else:
                if barcode in seen:
                    problems.append(("barcode", "duplicate_in_file"))
                seen.add(barcode)

        category = value(row, "category")
        if len(category) > 100:
            problems.append(("category", "too_long"))

        unit_text = value(row, "unit")
        unit = units.get(_key(unit_text)) if unit_text else default_unit
        if unit is None:
            problems.append(("unit", "unknown_unit"))

        prices = {}
        for key in ("sell_price", "cost_price"):
            text = value(row, key)
            if not text:
                if key == "sell_price":
                    problems.append((key, "required"))
                prices[key] = 0
                continue
            try:
                prices[key] = parse_money(text)
            except BadNumber:
                problems.append((key, "invalid_number"))
                continue
            if prices[key] < 0 or (key == "sell_price" and prices[key] == 0):
                problems.append((key, "not_positive"))

        stock = 0
        stock_text = value(row, "stock")
        if stock_text and unit is not None:
            try:
                stock = parse_qty(stock_text, fractional=unit.is_fractional)
            except BadNumber:
                problems.append(("stock", "invalid_number"))
            else:
                if stock < 0:
                    problems.append(("stock", "not_positive"))

        product = existing.get(barcode) if barcode else None
        if product is not None and product.is_weighted and unit and not unit.is_fractional:
            problems.append(("unit", "weighted_unit"))

        if problems:
            result.errors.extend({"row": number, "field": f, "code": c} for f, c in problems)
            continue
        result.rows.append(
            ParsedRow(
                row=number,
                name=name,
                barcode=barcode,
                category=category,
                unit=unit,
                sell=prices["sell_price"],
                cost=prices.get("cost_price", 0),
                stock=stock,
                product=product,
            )
        )
    return result


def _log(tenant_id, entity: str, ids) -> None:
    ChangeLog.objects.bulk_create(
        ChangeLog(tenant_id=tenant_id, entity=entity, entity_id=pk, op=ChangeOp.UPSERT)
        for pk in ids
    )


@transaction.atomic
def apply_import(product_import: ProductImport, mapping: dict[str, int | None], *, user) -> dict:
    """To'g'ri qatorlar bitta tranzaksiyada; xatolilari o'tkazib yuboriladi."""
    from apps.inventory.models import (
        DocumentStatus,
        MovementType,
        ProductStock,
        StockMovement,
        StockReceipt,
        StockReceiptItem,
    )
    from apps.inventory.services import default_store, next_receipt_number

    product_import = ProductImport.objects.select_for_update().get(pk=product_import.pk)
    if product_import.status == ProductImport.Status.APPLIED:
        raise DomainError("import_already_applied", "This import was already applied.", 409)
    tenant_id = product_import.tenant_id
    analysis = analyze(product_import, mapping)
    now = timezone.now()

    # Kategoriyalar: nomi bo'yicha, yo'qlari yaratiladi.
    categories = {
        _key(c.name): c
        for c in Category.objects.filter(tenant_id=tenant_id, deleted_at__isnull=True)
    }
    new_categories = []
    for row in analysis.rows:
        key = _key(row.category)
        if row.category and key not in categories:
            category = Category(id=new_id(), tenant_id=tenant_id, name=row.category)
            categories[key] = category
            new_categories.append(category)
    Category.objects.bulk_create(new_categories)
    _log(tenant_id, "category", [c.pk for c in new_categories])

    created, updated, history = [], [], []
    barcodes = []
    for row in analysis.rows:
        category = categories.get(_key(row.category)) if row.category else None
        product = row.product
        if product is None:
            product = Product(
                id=new_id(),
                tenant_id=tenant_id,
                name=row.name,
                category=category,
                unit=row.unit,
                sell_price_tiyin=row.sell,
                cost_price_tiyin=row.cost,
            )
            created.append((product, row))
            if row.barcode:
                barcodes.append(
                    ProductBarcode(
                        tenant_id=tenant_id, product=product, barcode=row.barcode, is_primary=True
                    )
                )
        else:
            old_prices = (product.sell_price_tiyin, product.cost_price_tiyin)
            product.name = row.name
            product.unit = row.unit
            if row.category:
                product.category = category
            product.sell_price_tiyin = row.sell
            if mapping.get("cost_price") is not None:
                product.cost_price_tiyin = row.cost
            product.updated_at = now
            updated.append(product)
            if (product.sell_price_tiyin, product.cost_price_tiyin) == old_prices:
                continue
        history.append(
            PriceHistory(
                tenant_id=tenant_id,
                product=product,
                sell_price_tiyin=product.sell_price_tiyin,
                cost_price_tiyin=product.cost_price_tiyin,
                changed_by=user,
            )
        )

    Product.objects.bulk_create([p for p, _ in created])
    if updated:
        Product.objects.bulk_update(
            updated,
            ["name", "unit", "category", "sell_price_tiyin", "cost_price_tiyin", "updated_at"],
            batch_size=1000,
        )
    ProductBarcode.objects.bulk_create(barcodes)
    PriceHistory.objects.bulk_create(history)
    _log(tenant_id, "product", [p.pk for p, _ in created] + [p.pk for p in updated])

    # Boshlang'ich qoldiq — faqat yangi tovarlar: bitta tasdiqlangan kirim.
    stocked = [(p, r) for p, r in created if r.stock > 0]
    receipt = None
    if stocked:
        store = default_store(tenant_id)
        items = [
            StockReceiptItem(
                tenant_id=tenant_id,
                product=p,
                qty_milli=r.stock,
                cost_price_tiyin=r.cost,
                line_total_tiyin=line_total(r.cost, r.stock),
            )
            for p, r in stocked
        ]
        receipt = StockReceipt.objects.create(
            tenant_id=tenant_id,
            store=store,
            number=next_receipt_number(tenant_id),
            status=DocumentStatus.POSTED,
            total_cost_tiyin=sum(i.line_total_tiyin for i in items),
            note="Import",
            created_by=user,
            posted_at=now,
            posted_by=user,
        )
        for item in items:
            item.receipt = receipt
        StockReceiptItem.objects.bulk_create(items)
        StockMovement.objects.bulk_create(
            StockMovement(
                tenant_id=tenant_id,
                product=p,
                store=store,
                qty_milli=r.stock,
                type=MovementType.RECEIPT,
                ref_type="stock_receipt",
                ref_id=receipt.pk,
                created_by=user,
                at=now,
            )
            for p, r in stocked
        )
        # Yangi tovarlarda kesh qatori yo'q — to'g'ridan-to'g'ri yaratiladi.
        ProductStock.objects.bulk_create(
            ProductStock(tenant_id=tenant_id, product=p, store=store, qty_milli=r.stock)
            for p, r in stocked
        )
        _log(tenant_id, "stock", [p.pk for p, _ in stocked])

    result = {
        "created": len(created),
        "updated": len(updated),
        "skipped": len({e["row"] for e in analysis.errors}),
        "error_count": len(analysis.errors),
        "errors": analysis.errors[:MAX_ERRORS],
        "receipt": receipt.number if receipt else None,
    }
    product_import.status = ProductImport.Status.APPLIED
    product_import.result = result
    product_import.save(update_fields=["status", "result", "updated_at"])
    return result
